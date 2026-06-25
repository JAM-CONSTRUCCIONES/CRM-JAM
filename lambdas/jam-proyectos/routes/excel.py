import json
import uuid
import boto3
import os
import re
from io import BytesIO
from datetime import datetime
from boto3.dynamodb.conditions import Key
from utils.response import ok, bad_request, not_found
from utils.helpers import now
from utils.excel_parser import parsear_excel  # noqa: F401 — usado por jam-excel-parser
from utils.excel_writer import escribir_migracion

s3_client    = boto3.client('s3')
dynamodb     = boto3.resource('dynamodb')
inventario_t = dynamodb.Table(os.environ['INVENTARIO_TABLE'])
clientes_t   = dynamodb.Table(os.environ.get('CLIENTES_TABLE', ''))
procesos_t   = dynamodb.Table(os.environ.get('PROCESOS_TABLE', ''))
usuarios_t   = dynamodb.Table(os.environ.get('USUARIOS_TABLE', ''))

ARCHIVOS_BUCKET = os.environ.get('ARCHIVOS_BUCKET', '')
EMAIL_RE = re.compile(r'^[^\s@]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}$')
DOCUMENTO_RE = re.compile(r'^[A-Z0-9]+$')


def _clean_spaces(value):
    return ' '.join(str(value or '').strip().split())


def _nombre_valido(value):
    nombre = _clean_spaces(value)
    return bool(nombre) and all(ch.isalpha() or ch.isspace() for ch in nombre)


def get_upload_url(event):
    qs          = event.get('queryStringParameters') or {}
    proyecto_id = qs.get('proyecto_id', '').strip()
    if not proyecto_id:
        return bad_request('proyecto_id requerido')

    job_id = str(uuid.uuid4())
    key    = f'inventario/uploads/{proyecto_id}/{job_id}.xlsx'

    url = s3_client.generate_presigned_url(
        'put_object',
        Params={
            'Bucket':      ARCHIVOS_BUCKET,
            'Key':         key,
            'ContentType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'Metadata':    {'proyecto_id': proyecto_id, 'job_id': job_id},
        },
        ExpiresIn=300,
    )
    return ok({'upload_url': url, 'job_id': job_id, 'key': key})


def get_preview(job_id):
    try:
        obj = s3_client.get_object(Bucket=ARCHIVOS_BUCKET, Key=f'inventario/previews/{job_id}.json')
        return ok(json.loads(obj['Body'].read()))
    except Exception:
        return not_found('Preview no encontrado o expirado')


def confirmar(job_id, event):
    body           = json.loads(event.get('body') or '{}')
    filas_editadas = body.get('filas', [])

    try:
        obj     = s3_client.get_object(Bucket=ARCHIVOS_BUCKET, Key=f'inventario/previews/{job_id}.json')
        preview = json.loads(obj['Body'].read())
    except Exception:
        return not_found('Preview no encontrado o expirado')

    proyecto_id = preview['proyecto_id']

    filas = [dict(f) for f in preview['filas']]
    editadas_por_fila = {
        fe.get('fila_excel'): fe
        for fe in filas_editadas
        if fe.get('fila_excel') is not None
    }
    editadas_por_unidad = {
        fe.get('id_unidad'): fe
        for fe in filas_editadas
        if fe.get('id_unidad')
    }

    for fila in filas:
        fila_excel = fila.get('fila_excel')
        id_unidad = fila.get('id_unidad')
        if fila_excel in editadas_por_fila:
            fila.update(editadas_por_fila[fila_excel])
        elif id_unidad in editadas_por_unidad:
            fila.update(editadas_por_unidad[id_unidad])

    errores_duplicados = _validar_duplicados_confirmacion(proyecto_id, filas)
    if errores_duplicados:
        return bad_request(
            f'{len(errores_duplicados)} unidad(es) ya existen o estan duplicadas: ' +
            ', '.join(errores_duplicados[:10])
        )

    errores_clientes = _validar_clientes_confirmacion(filas)
    if errores_clientes:
        return bad_request(
            f'{len(errores_clientes)} fila(s) tienen datos de cliente invalidos: ' +
            ', '.join(errores_clientes[:10])
        )

    errores_estatus = _validar_estatus_confirmacion(filas)
    if errores_estatus:
        return bad_request(
            f'{len(errores_estatus)} fila(s) tienen estado y estatus de proceso incompatibles: ' +
            ', '.join(errores_estatus[:10])
        )

    filas_con_error = [f for f in filas if f.get('errores')]
    if filas_con_error:
        return bad_request(f'{len(filas_con_error)} filas tienen errores. Corrigelas antes de confirmar.')

    try:
        stats = escribir_migracion(proyecto_id, filas)
    except ValueError as e:
        return bad_request(str(e))

    reporte = {
        'job_id':        job_id,
        'archivo':       preview.get('archivo', ''),
        'proyecto_id':   proyecto_id,
        'total_filas':   len(filas),
        **stats,
        'completado_en': now(),
    }
    s3_client.put_object(
        Bucket=ARCHIVOS_BUCKET,
        Key=f'inventario/reportes/{job_id}.json',
        Body=json.dumps(reporte, ensure_ascii=False, default=str),
        ContentType='application/json',
    )

    # Borrar el Excel original — ya no se necesita después de confirmar
    try:
        excel_key = f'inventario/uploads/{proyecto_id}/{job_id}.xlsx'
        s3_client.delete_object(Bucket=ARCHIVOS_BUCKET, Key=excel_key)
    except Exception:
        pass  # No crítico si falla el borrado

    return ok(reporte)


def _normalizar_id_unidad(id_unidad):
    return str(id_unidad or '').strip().upper()


def _validar_duplicados_confirmacion(proyecto_id, filas):
    errores = []
    vistos = set()
    existentes = _obtener_ids_unidades_existentes(proyecto_id)

    for fila in filas:
        id_unidad = (fila.get('id_unidad') or '').strip()
        id_norm = _normalizar_id_unidad(id_unidad)
        if not id_norm:
            continue

        if id_norm in vistos:
            errores.append(f'{id_unidad} duplicada en el Excel')
            continue
        vistos.add(id_norm)

        if id_norm in existentes:
            errores.append(f'{id_unidad} ya existe en el proyecto')

    return errores


def _validar_clientes_confirmacion(filas):
    errores = []
    for fila in filas:
        cedula = str(fila.get('cedula') or '').strip().upper()
        nombres = _clean_spaces(fila.get('nombres_sugerido'))
        apellidos = _clean_spaces(fila.get('apellidos_sugerido'))
        telefono = _clean_spaces(fila.get('cliente_telefono'))
        correo = _clean_spaces(fila.get('cliente_correo')).lower()
        tiene_datos_cliente = any(_clean_spaces(fila.get(campo)) for campo in (
            'cliente_nombre_excel',
            'cedula',
            'nombres_sugerido',
            'apellidos_sugerido',
            'cliente_telefono',
            'cliente_correo',
            'fecha_nacimiento',
            'pais_residencia',
            'nacionalidad',
        ))
        fila_excel = fila.get('fila_excel', '?')
        id_unidad = fila.get('id_unidad', '')

        if tiene_datos_cliente and not cedula:
            errores.append(f'fila {fila_excel} unidad {id_unidad}: documento de identidad obligatorio')
        if cedula and not DOCUMENTO_RE.match(cedula):
            errores.append(f'fila {fila_excel} unidad {id_unidad}: documento solo debe contener letras y numeros')
        if tiene_datos_cliente and not nombres:
            errores.append(f'fila {fila_excel} unidad {id_unidad}: nombres requeridos')
        elif nombres and not _nombre_valido(nombres):
            errores.append(f'fila {fila_excel} unidad {id_unidad}: nombres solo deben contener letras y espacios')
        if tiene_datos_cliente and not apellidos:
            errores.append(f'fila {fila_excel} unidad {id_unidad}: apellidos requeridos')
        elif apellidos and not _nombre_valido(apellidos):
            errores.append(f'fila {fila_excel} unidad {id_unidad}: apellidos solo deben contener letras y espacios')
        if telefono and not telefono.isdigit():
            errores.append(f'fila {fila_excel} unidad {id_unidad}: telefono solo debe contener numeros')
        if correo and not EMAIL_RE.match(correo):
            errores.append(f'fila {fila_excel} unidad {id_unidad}: correo invalido')
        for campo, etiqueta in (
            ('fecha_nacimiento', 'fecha de nacimiento'),
            ('fecha_bloqueo', 'Fecha Bloqueo'),
        ):
            valor = str(fila.get(campo) or '').strip()
            if valor and not _fecha_valida(valor):
                errores.append(f'fila {fila_excel} unidad {id_unidad}: {etiqueta} invalida')
    return errores


def _validar_estatus_confirmacion(filas):
    errores = []
    for fila in filas:
        estado = str(fila.get('estado') or '').strip()
        estado_proceso = str(fila.get('estado_proceso') or '').strip()
        fila_excel = fila.get('fila_excel', '?')
        id_unidad = fila.get('id_unidad', '')
        mensaje = _validar_estado_proceso(estado, estado_proceso)
        if mensaje:
            errores.append(f'fila {fila_excel} unidad {id_unidad}: {mensaje}')
    return errores


def _validar_estado_proceso(estado, estado_proceso):
    permitidos = {
        'disponible': {''},
        'bloqueada': {'captacion'},
        'no_disponible': {'reserva', 'separacion', 'inicial', 'pagos_atrasados', 'contra_entrega'},
        'vendida': {'vendido'},
    }
    if estado not in permitidos:
        return f'estado invalido: {estado}'
    if estado_proceso not in permitidos[estado]:
        etiquetas = {
            'disponible': 'sin estatus de proceso',
            'bloqueada': 'captacion',
            'no_disponible': 'reserva, separacion, contrato firmado, pagos atrasados o contra entrega',
            'vendida': 'vendido',
        }
        return f'{estado} debe usar {etiquetas[estado]}'
    return None


def _fecha_valida(valor):
    try:
        datetime.fromisoformat(valor.replace('Z', '+00:00'))
        return True
    except ValueError:
        return False


def _obtener_ids_unidades_existentes(proyecto_id):
    ids = set()
    last_key = None

    while True:
        kwargs = {
            'ProjectionExpression': 'id_unidad',
            'KeyConditionExpression': Key('pk').eq(f'PROYECTO#{proyecto_id}') & Key('sk').begins_with('UNIDAD#'),
        }
        if last_key:
            kwargs['ExclusiveStartKey'] = last_key

        result = inventario_t.query(**kwargs)
        for item in result.get('Items', []):
            id_unidad = _normalizar_id_unidad(item.get('id_unidad'))
            if id_unidad:
                ids.add(id_unidad)

        last_key = result.get('LastEvaluatedKey')
        if not last_key:
            return ids


def listar_reportes(event):
    """GET /admin/inventario/reportes?proyecto_id=xxx — lista historial de migraciones."""
    qs          = event.get('queryStringParameters') or {}
    proyecto_id = qs.get('proyecto_id', '').strip()

    paginator = s3_client.get_paginator('list_objects_v2')
    pages     = paginator.paginate(Bucket=ARCHIVOS_BUCKET, Prefix='inventario/reportes/')

    reportes = []
    for page in pages:
        for obj in page.get('Contents', []):
            try:
                body = s3_client.get_object(Bucket=ARCHIVOS_BUCKET, Key=obj['Key'])
                data = json.loads(body['Body'].read())
                # Filtrar por proyecto si se especifica
                if proyecto_id and data.get('proyecto_id') != proyecto_id:
                    continue
                reportes.append({
                    'job_id':           data.get('job_id'),
                    'archivo':          data.get('archivo'),
                    'proyecto_id':      data.get('proyecto_id'),
                    'total_filas':      data.get('total_filas', 0),
                    'unidades_cargadas': data.get('unidades_cargadas', 0),
                    'clientes_creados': data.get('clientes_creados', 0),
                    'inmobiliarias_creadas': data.get('inmobiliarias_creadas', 0),
                    'advertencias':     data.get('advertencias', 0),
                    'completado_en':    data.get('completado_en'),
                })
            except Exception:
                continue

    # Ordenar por fecha descendente
    reportes.sort(key=lambda r: r.get('completado_en') or '', reverse=True)
    return ok(reportes)


def get_reporte(job_id):
    try:
        obj = s3_client.get_object(Bucket=ARCHIVOS_BUCKET, Key=f'inventario/reportes/{job_id}.json')
        return ok(json.loads(obj['Body'].read()))
    except Exception:
        return not_found('Reporte no encontrado')



def exportar_inventario(event):
    """GET /admin/inventario/exportar?proyecto_id=xxx — exporta inventario completo a Excel."""
    qs = event.get('queryStringParameters') or {}
    proyecto_id = qs.get('proyecto_id', '').strip()
    
    if not proyecto_id:
        return bad_request('proyecto_id requerido')
    
    try:
        # Importar openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from decimal import Decimal
        
        # 1. Obtener proyecto
        proyecto = inventario_t.get_item(
            Key={'pk': f'PROYECTO#{proyecto_id}', 'sk': 'METADATA'}
        ).get('Item')
        
        if not proyecto:
            return not_found('Proyecto no encontrado')
        
        proyecto_nombre = proyecto.get('nombre', proyecto_id)
        
        # 2. Obtener todas las unidades del proyecto
        unidades = []
        last_key = None
        
        while True:
            query_kwargs = {
                'KeyConditionExpression': Key('pk').eq(f'PROYECTO#{proyecto_id}') & Key('sk').begins_with('UNIDAD#'),
                'Limit': 500
            }
            if last_key:
                query_kwargs['ExclusiveStartKey'] = last_key
            
            result = inventario_t.query(**query_kwargs)
            unidades.extend(result.get('Items', []))
            
            last_key = result.get('LastEvaluatedKey')
            if not last_key:
                break
        
        if not unidades:
            return bad_request('El proyecto no tiene unidades para exportar')
        
        # 3. Obtener etapas para nombres
        etapas_result = inventario_t.query(
            KeyConditionExpression=Key('pk').eq(f'PROYECTO#{proyecto_id}') & Key('sk').begins_with('ETAPA#')
        )
        etapas_map = {e['etapa_id']: e.get('nombre', e['etapa_id']) for e in etapas_result.get('Items', [])}
        
        # 4. Obtener procesos del proyecto
        procesos_map = {}
        try:
            procesos_result = procesos_t.query(
                IndexName='gsi-proyecto-procesos',
                KeyConditionExpression=Key('proyecto_id').eq(proyecto_id)
            )
            for p in procesos_result.get('Items', []):
                unidad_id = p.get('unidad_id', '')
                if unidad_id:
                    procesos_map[unidad_id] = p
        except Exception as e:
            print(f'Error al cargar procesos: {e}')
        
        # 5. Obtener clientes
        clientes_map = {}
        try:
            for proceso in procesos_map.values():
                cedula = proceso.get('cedula')
                inmo_id = proceso.get('inmobiliaria_id')
                if cedula and inmo_id:
                    # Intentar con formato corto primero
                    pk_corto = f'CLIENTE#{cedula}#{inmo_id}'
                    sk = f'PROYECTO#{proyecto_id}'
                    
                    cliente = clientes_t.get_item(Key={'pk': pk_corto, 'sk': sk}).get('Item')
                    
                    # Si no existe, intentar con formato largo
                    if not cliente:
                        pk_largo = f'CLIENTE#{cedula}#INMOBILIARIA#{inmo_id}'
                        cliente = clientes_t.get_item(Key={'pk': pk_largo, 'sk': sk}).get('Item')
                    
                    if cliente:
                        clientes_map[f'{cedula}#{inmo_id}'] = cliente
        except Exception as e:
            print(f'Error al cargar clientes: {e}')
            import traceback
            traceback.print_exc()
        
        # 6. Obtener inmobiliarias
        inmobiliarias_map = {}
        try:
            inmos_result = usuarios_t.query(
                IndexName='gsi-tipo',
                KeyConditionExpression=Key('tipo').eq('INMOBILIARIA')
            )
            for inmo in inmos_result.get('Items', []):
                inmo_id = inmo['pk'].replace('INMOBILIARIA#', '')
                inmobiliarias_map[inmo_id] = inmo.get('nombre', inmo_id)
        except Exception as e:
            print(f'Error al cargar inmobiliarias: {e}')
        
        # 7. Crear Excel
        wb = Workbook()
        ws = wb.active
        ws.title = 'Inventario'
        
        # Estilos
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=11)
        border = Border(
            left=Side(style='thin', color='D0D0D0'),
            right=Side(style='thin', color='D0D0D0'),
            top=Side(style='thin', color='D0D0D0'),
            bottom=Side(style='thin', color='D0D0D0')
        )
        
        # Headers
        headers = [
            'ID Unidad', 'Etapa', 'Tipo', 'Bloque', 'Piso',
            'Metraje', 'Metraje Terraza', 'Metraje Patio',
            'Cuartos', 'Baños', 'Parqueos',
            'Precio', 'Precio Reserva', 'Precio Separación', 'Precio Inicial',
            'Cuota Monto', 'Cuota Meses', 'Contra Entrega',
            'Estado', 'Documento', 'Cliente Nombre', 'Cliente Teléfono', 'Cliente Correo',
            'Fecha Nacimiento', 'Pais Residencia', 'Nacionalidad',
            'Inmobiliaria', 'Estado Proceso',
            'Fecha Bloqueo', 'Fecha Liberación', 'Días Restantes',
            'Comentario'
        ]
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Ajustar anchos de columna
        column_widths = [15, 15, 12, 10, 8, 10, 12, 12, 8, 8, 10, 12, 12, 12, 12, 12, 10, 12, 14, 15, 25, 15, 25, 18, 18, 18, 20, 15, 18, 18, 12, 30]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        
        # Datos
        row_num = 2
        for unidad in sorted(unidades, key=lambda u: u.get('id_unidad', '')):
            unidad_id = unidad.get('unidad_id', '')
            proceso = procesos_map.get(unidad_id, {})
            
            cedula = proceso.get('cedula', '')
            inmo_id = proceso.get('inmobiliaria_id', '')
            cliente = clientes_map.get(f'{cedula}#{inmo_id}', {})
            
            # Calcular días restantes
            dias_restantes = ''
            if unidad.get('estado') == 'bloqueada' and unidad.get('fecha_liberacion'):
                try:
                    from datetime import datetime, timezone
                    dt_lib = datetime.fromisoformat(unidad['fecha_liberacion'].replace('Z', '+00:00'))
                    restante = int((dt_lib - datetime.now(timezone.utc)).total_seconds() / 86400)
                    dias_restantes = max(restante, 0)
                except Exception:
                    pass
            
            def fmt(val):
                """Formatear valores Decimal y None."""
                if val is None or val == '':
                    return ''
                if isinstance(val, Decimal):
                    return float(val)
                return val

            def estado_export(val):
                estados = {
                    'disponible': 'DISPONIBLE',
                    'bloqueada': 'BLOQUEADO',
                    'no_disponible': 'NO DISPONIBLE',
                }
                return estados.get(str(val or '').strip().lower(), val or '')
            
            row_data = [
                unidad.get('id_unidad', ''),
                etapas_map.get(unidad.get('etapa_id', ''), unidad.get('etapa_id', '')),
                unidad.get('tipo', ''),
                unidad.get('manzana', ''),
                unidad.get('piso', ''),
                fmt(unidad.get('metraje')),
                fmt(unidad.get('metraje_terraza')),
                fmt(unidad.get('metraje_patio')),
                fmt(unidad.get('num_cuartos')),
                fmt(unidad.get('num_banos')),
                fmt(unidad.get('parqueos')),
                fmt(unidad.get('precio')),
                fmt(unidad.get('precio_reserva')),
                fmt(unidad.get('precio_separacion')),
                fmt(unidad.get('precio_inicial')),
                fmt(unidad.get('cuota_monto')),
                fmt(unidad.get('cuota_meses')),
                fmt(unidad.get('contra_entrega')),
                estado_export(unidad.get('estado', '')),
                cliente.get('cedula', ''),
                f"{cliente.get('nombres', '')} {cliente.get('apellidos', '')}".strip(),
                cliente.get('telefono', ''),
                cliente.get('correo', ''),
                cliente.get('fecha_nacimiento', ''),
                cliente.get('pais_residencia', ''),
                cliente.get('nacionalidad', ''),
                inmobiliarias_map.get(inmo_id, ''),
                proceso.get('estado', ''),
                unidad.get('fecha_bloqueo', ''),
                unidad.get('fecha_liberacion', ''),
                dias_restantes,
                unidad.get('comentario', '')
            ]
            
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical='center')
            
            row_num += 1
        
        # Guardar en memoria
        excel_buffer = BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        # Subir a S3
        timestamp = now().replace(':', '-').replace('.', '-')
        filename = f'inventario_{proyecto_id}_{timestamp}.xlsx'
        s3_key = f'inventario/exportaciones/{filename}'
        
        s3_client.put_object(
            Bucket=ARCHIVOS_BUCKET,
            Key=s3_key,
            Body=excel_buffer.getvalue(),
            ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            ContentDisposition=f'attachment; filename="{filename}"'
        )
        
        # Generar URL de descarga (válida por 5 minutos)
        download_url = s3_client.generate_presigned_url(
            'get_object',
            Params={'Bucket': ARCHIVOS_BUCKET, 'Key': s3_key},
            ExpiresIn=300
        )
        
        return ok({
            'download_url': download_url,
            'filename': filename,
            'total_unidades': len(unidades),
            'proyecto_nombre': proyecto_nombre
        })
        
    except Exception as e:
        print(f'ERROR exportar_inventario: {type(e).__name__}: {e}')
        import traceback
        traceback.print_exc()
        return bad_request(f'Error al generar exportación: {str(e)}')
